import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from calculos import calcular_modelo, calcular_multas_7317
from excel_io import (
    cargar_libro_referencia,
    crear_reporte_xlsx,
    fila_costo_desde_referencia,
    fila_receptor_desde_caso,
)


class ReferenciasVdepTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        libro = Path(__file__).resolve().parents[1] / "data" / "Plantilla_Calculadora_ESVD_ES.xlsx"
        (
            cls.config,
            cls.parametros,
            cls.especies,
            cls.costos,
            cls.casos,
            cls.receptores,
            cls.variables,
            cls.multas,
        ) = cargar_libro_referencia(libro)

    def calcular_caso(self, id_caso):
        costos = [
            fila_costo_desde_referencia(c)
            for c in self.costos
            if c["ID_CASO"] == id_caso
        ]
        return calcular_modelo([], costos, self.config), costos

    def test_libro_contiene_nueve_casos_aplicados(self):
        self.assertEqual(self.config["tipo_valoracion"], "VDEP")
        self.assertEqual(self.config["moneda_modelo"], "CRC")
        self.assertEqual(len(self.casos), 9)
        self.assertEqual(len(self.receptores), 11)
        self.assertEqual(len(self.variables), 49)
        self.assertEqual(len(self.costos), 32)
        self.assertEqual(len(self.parametros), 4)
        self.assertEqual(len(self.multas), 35)
        self.assertEqual(self.config["salario_base_ley_7337"], 462200)
        self.assertTrue(all(str(c["ID_CASO"]).startswith("VDEP-EJ") for c in self.casos))

    def test_los_nueve_totales_reproducen_el_anexo(self):
        esperados = {
            "VDEP-EJ01": (400000, 470000, 580000),
            "VDEP-EJ02": (31040000, 36501000, 41758000),
            "VDEP-EJ03": (9520000, 12120000, 15845000),
            "VDEP-EJ04": (5200000, 7100000, 9200000),
            "VDEP-EJ05": (8970000, 10650000, 12750000),
            "VDEP-EJ06": (2950000, 3500000, 4050000),
            "VDEP-EJ07": (6075000, 8450000, 11025000),
            "VDEP-EJ08": (3750000, 4550000, 5350000),
            "VDEP-EJ09": (12580000, 15200000, 18470000),
        }
        for id_caso, esperado in esperados.items():
            with self.subTest(id_caso=id_caso):
                resultado, _ = self.calcular_caso(id_caso)
                obtenido = tuple(resultado["total_compatible"][e] for e in ("bajo", "central", "alto"))
                self.assertEqual(obtenido, esperado)
                caso = next(c for c in self.casos if c["ID_CASO"] == id_caso)
                publicado = (caso["TOTAL_BAJO"], caso["TOTAL_CENTRAL"], caso["TOTAL_ALTO"])
                self.assertEqual(obtenido, publicado)

    def test_ejemplo_uno_conserva_r_separada_y_total_compatible(self):
        resultado, costos = self.calcular_caso("VDEP-EJ01")
        self.assertEqual(resultado["principal_por_cuenta"]["C"], {"bajo": 210000, "central": 210000, "alto": 210000})
        self.assertEqual(resultado["principal_por_cuenta"]["D"], {"bajo": 100000, "central": 140000, "alto": 210000})
        self.assertEqual(resultado["reparacion"], {"bajo": 90000, "central": 120000, "alto": 160000})
        self.assertEqual(resultado["subtotal_danos"], {"bajo": 310000, "central": 350000, "alto": 420000})
        self.assertEqual(resultado["total_compatible"], {"bajo": 400000, "central": 470000, "alto": 580000})
        d = next(c for c in costos if c["id_referencia"] == "VDEP-EJ01-D01")
        self.assertEqual(d["n_referencia"], "42")

    def test_receptores_y_reporte_exportable(self):
        receptores = [
            fila_receptor_desde_caso(r)
            for r in self.receptores
            if r["ID_CASO"] == "VDEP-EJ09"
        ]
        self.assertEqual(len(receptores), 2)
        self.assertEqual(receptores[1]["cantidad"], 0.6)
        resultado, costos = self.calcular_caso("VDEP-EJ09")
        variables = [v for v in self.variables if v["ID_CASO"] == "VDEP-EJ09"]
        multas = calcular_multas_7317(
            [{
                "aplica": True,
                "articulo": "90",
                "conducta": "Prueba de exportación",
                "min_sb": 1,
                "sb_aplicados": 2,
                "max_sb": 3,
            }],
            462200,
            1,
        )
        reporte = crear_reporte_xlsx(
            {"moneda": "CRC", "permitir_agregacion": True, "tipo_valoracion": "VDEP"},
            resultado,
            receptores,
            [],
            costos,
            variables,
            multas,
        )
        libro = load_workbook(BytesIO(reporte), read_only=True, data_only=False)
        self.assertIn("RESUMEN", libro.sheetnames)
        self.assertIn("VARIABLES_BIOFISICAS", libro.sheetnames)
        self.assertIn("MULTAS_LEY_7317", libro.sheetnames)
        valores_columna_a = [libro["RESUMEN"].cell(r, 1).value for r in range(1, libro["RESUMEN"].max_row + 1)]
        self.assertIn("TOTAL VDEP", valores_columna_a)
        self.assertIn("MULTAS LEY 7317", valores_columna_a)
        self.assertIn("TOTAL FINAL", valores_columna_a)
        libro.close()


if __name__ == "__main__":
    unittest.main()
