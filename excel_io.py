"""Lectura del libro de referencia y exportación del informe de resultados."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from calculos import ESCENARIOS, normalizar_valores_esvd, numero, texto


HOJA_VALORES = "VALORES_ESVD"
HOJA_CONFIG = "CONFIGURACION"
HOJA_ESPECIES = "ESPECIES_REFERENCIA"
HOJA_COSTOS = "COSTOS_REFERENCIA"


def _abrir(origen: str | Path | bytes | BinaryIO):
    if isinstance(origen, bytes):
        origen = BytesIO(origen)
    return load_workbook(origen, read_only=True, data_only=False)


def _leer_tabla(ws, encabezado_fila: int, clave: str) -> list[dict[str, Any]]:
    encabezados = [texto(c.value) for c in ws[encabezado_fila]]
    if clave not in encabezados:
        raise ValueError(f"La hoja {ws.title} no contiene el encabezado {clave} en la fila {encabezado_fila}.")
    filas: list[dict[str, Any]] = []
    for valores in ws.iter_rows(min_row=encabezado_fila + 1, values_only=True):
        fila = dict(zip(encabezados, valores))
        if texto(fila.get(clave)):
            filas.append(fila)
    return filas


def cargar_libro_referencia(
    origen: str | Path | bytes | BinaryIO,
) -> tuple[dict[str, Any], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    libro = _abrir(origen)
    if HOJA_CONFIG not in libro.sheetnames or HOJA_VALORES not in libro.sheetnames:
        raise ValueError("El archivo debe contener las hojas CONFIGURACION y VALORES_ESVD.")

    ws_cfg = libro[HOJA_CONFIG]
    config: dict[str, Any] = {}
    for fila in ws_cfg.iter_rows(min_row=5, values_only=True):
        clave = texto(fila[0])
        if clave:
            config[clave] = fila[1]

    parametros = [normalizar_valores_esvd(fila) for fila in _leer_tabla(libro[HOJA_VALORES], 5, "ID_PARAMETRO")]
    especies = _leer_tabla(libro[HOJA_ESPECIES], 5, "ID_REFERENCIA") if HOJA_ESPECIES in libro.sheetnames else []
    costos = _leer_tabla(libro[HOJA_COSTOS], 5, "ID_REFERENCIA") if HOJA_COSTOS in libro.sheetnames else []
    libro.close()
    return config, parametros, especies, costos


def fila_servicio_desde_parametro(parametro: dict[str, Any]) -> dict[str, Any]:
    """Convierte una referencia ESVD en una fila editable del formulario."""
    return {
        "id_parametro": texto(parametro.get("ID_PARAMETRO")),
        "servicio": texto(parametro.get("SERVICIO_ECOSISTEMICO")),
        "unidad_base": texto(parametro.get("UNIDAD_BASE")),
        "unidades_afectadas": numero(parametro.get("UNIDADES_AFECTADAS_DEMO"), 0.0),
        "valor_unitario_bajo": numero(parametro.get("VALOR_BAJO_NORMALIZADO")),
        "valor_unitario_central": numero(parametro.get("VALOR_CENTRAL_NORMALIZADO")),
        "valor_unitario_alto": numero(parametro.get("VALOR_ALTO_NORMALIZADO")),
        "perdida_inicial_bajo": numero(parametro.get("PERDIDA_INICIAL_BAJA_DEMO"), 0.0),
        "perdida_inicial_central": numero(parametro.get("PERDIDA_INICIAL_CENTRAL_DEMO"), 0.0),
        "perdida_inicial_alto": numero(parametro.get("PERDIDA_INICIAL_ALTA_DEMO"), 0.0),
        "recuperacion_bajo": int(numero(parametro.get("RECUPERACION_BAJA_DEMO"), 1)),
        "recuperacion_central": int(numero(parametro.get("RECUPERACION_CENTRAL_DEMO"), 1)),
        "recuperacion_alto": int(numero(parametro.get("RECUPERACION_ALTA_DEMO"), 1)),
        "perfil_recuperacion": "Lineal",
        "evidencia": "B",
        "evidencia_parametro": texto(parametro.get("NIVEL_EVIDENCIA")),
        "comparabilidad": texto(parametro.get("COMPARABILIDAD")) or "Media",
        "nexo_causal": "Sí",
        "beneficiarios": texto(parametro.get("BENEFICIARIOS")),
        "fuente": texto(parametro.get("FUENTE_URL_CITA")),
        "grupo_doble_conteo": "",
    }


def fila_costo_desde_referencia(referencia: dict[str, Any]) -> dict[str, Any]:
    """Convierte un costo de referencia en una fila editable del formulario."""
    return {
        "id_referencia": texto(referencia.get("ID_REFERENCIA")),
        "cuenta": texto(referencia.get("CUENTA")) or "E",
        "concepto": texto(referencia.get("CONCEPTO")),
        "unidad": texto(referencia.get("UNIDAD")),
        "anio_desde_evento": int(numero(referencia.get("ANIO_DESDE_EVENTO"), 0)),
        "cantidad_bajo": numero(referencia.get("CANTIDAD_BAJA_DEMO"), 1.0),
        "cantidad_central": numero(referencia.get("CANTIDAD_CENTRAL_DEMO"), 1.0),
        "cantidad_alto": numero(referencia.get("CANTIDAD_ALTA_DEMO"), 1.0),
        "costo_unitario_bajo": numero(referencia.get("COSTO_UNITARIO_BAJO_USD")),
        "costo_unitario_central": numero(referencia.get("COSTO_UNITARIO_CENTRAL_USD")),
        "costo_unitario_alto": numero(referencia.get("COSTO_UNITARIO_ALTO_USD")),
        "evidencia": texto(referencia.get("NIVEL_EVIDENCIA")) or "B",
        "nexo_causal": texto(referencia.get("NEXO_CAUSAL")) or "Sí",
        "fuente": texto(referencia.get("FUENTE_DEMO")),
        "grupo_doble_conteo": texto(referencia.get("GRUPO_DOBLE_CONTEO")),
    }


def filas_especie_desde_referencia(
    referencia: dict[str, Any], cantidad: float,
) -> tuple[dict[str, Any], dict[str, Any]]:
    """Crea el registro biológico y su línea monetaria ficticia de Cuenta A."""
    cantidad = max(0.0, numero(cantidad))
    registro = {
        "id_referencia": texto(referencia.get("ID_REFERENCIA")),
        "especie_grupo": texto(referencia.get("ESPECIE_GRUPO")),
        "cantidad": cantidad,
        "tipo_afectacion": texto(referencia.get("TIPO_AFECTACION")),
        "funcion_ecologica": texto(referencia.get("FUNCION_ECOLOGICA")),
        "nexo_causal": texto(referencia.get("NEXO_CAUSAL")) or "Sí",
        "evidencia": texto(referencia.get("NIVEL_EVIDENCIA")) or "B",
        "fuente_expediente": texto(referencia.get("FUENTE_DEMO")),
        "notas": texto(referencia.get("NOTAS")),
    }
    costo = {
        "id_referencia": texto(referencia.get("ID_REFERENCIA")),
        "cuenta": "A",
        "concepto": f"{texto(referencia.get('ESPECIE_GRUPO'))}: {texto(referencia.get('TIPO_AFECTACION'))}",
        "unidad": texto(referencia.get("UNIDAD")),
        "anio_desde_evento": 0,
        "cantidad_bajo": cantidad,
        "cantidad_central": cantidad,
        "cantidad_alto": cantidad,
        "costo_unitario_bajo": numero(referencia.get("VALOR_BAJO_USD")),
        "costo_unitario_central": numero(referencia.get("VALOR_CENTRAL_USD")),
        "costo_unitario_alto": numero(referencia.get("VALOR_ALTO_USD")),
        "evidencia": texto(referencia.get("NIVEL_EVIDENCIA")) or "B",
        "nexo_causal": texto(referencia.get("NEXO_CAUSAL")) or "Sí",
        "fuente": texto(referencia.get("FUENTE_DEMO")),
        "grupo_doble_conteo": texto(referencia.get("GRUPO_DOBLE_CONTEO")),
    }
    return registro, costo


def _limpiar(valor: Any) -> Any:
    if valor is None:
        return ""
    if isinstance(valor, (str, int, float, bool, datetime)):
        return valor
    return str(valor)


def _ajustar_hoja(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for celda in ws[1]:
        celda.fill = PatternFill("solid", fgColor="147D80")
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(wrap_text=True, vertical="center")
    for col in range(1, ws.max_column + 1):
        largo = max((len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 100) + 1)), default=8)
        ws.column_dimensions[get_column_letter(col)].width = min(42, max(11, largo + 2))
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")


def crear_reporte_xlsx(
    caso: dict[str, Any],
    resultado: dict[str, Any],
    vida_silvestre: list[dict[str, Any]],
    servicios: list[dict[str, Any]],
    costos: list[dict[str, Any]],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "RESUMEN"
    moneda = texto(caso.get("moneda")) or "moneda del modelo"
    ws.append(["Cuenta", "Descripción", "Bajo", "Central", "Alto", "Tratamiento"])
    nombres = {
        "A": "Daño biológico monetizable",
        "B": "Pérdida de servicios ecosistémicos",
        "C": "Respuesta pública incremental",
        "D": "Rescate, rehabilitación y cuidado",
        "E": "Efectos conexos demostrados",
        "R": "Reparación/restauración/equivalencia",
    }
    for cuenta in ("A", "B", "C", "D", "E", "R"):
        valores = resultado["principal_por_cuenta"][cuenta]
        tratamiento = "Separada del subtotal de daños" if cuenta == "R" else "Cuenta principal"
        ws.append([cuenta, nombres[cuenta], *(valores[e] for e in ESCENARIOS), tratamiento])
    if caso.get("permitir_agregacion"):
        ws.append(["SUBTOTAL", f"Subtotal de daños ({moneda})", *(resultado["subtotal_danos"][e] for e in ESCENARIOS), "Agregación autorizada"])
    else:
        ws.append(["SUBTOTAL", "No mostrado: falta autorización de agregación", "", "", "", "Revise doble conteo"])
    ws.append(["EXPLORATORIO", "Total de líneas exploratorias", *(resultado["total_exploratorio"][e] for e in ESCENARIOS), "No integrar al escenario central"])
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = '#,##0.00'
    _ajustar_hoja(ws)

    hojas = [
        ("DETALLE_CALCULO", resultado["detalle"]),
        ("VIDA_SILVESTRE", vida_silvestre),
        ("SERVICIOS_ENTRADA", servicios),
        ("COSTOS_ENTRADA", costos),
    ]
    for nombre, filas in hojas:
        tab = wb.create_sheet(nombre)
        if filas:
            encabezados = list(filas[0].keys())
            tab.append(encabezados)
            for fila in filas:
                tab.append([_limpiar(fila.get(c)) for c in encabezados])
        else:
            tab.append(["Sin registros"])
        _ajustar_hoja(tab)

    adv = wb.create_sheet("ADVERTENCIAS")
    adv.append(["Advertencia"])
    for mensaje in resultado["advertencias"] or ["Sin advertencias automáticas."]:
        adv.append([mensaje])
    _ajustar_hoja(adv)

    meta = wb.create_sheet("METADATOS")
    meta.append(["Campo", "Valor"])
    for clave, valor in caso.items():
        meta.append([clave, _limpiar(valor)])
    meta.append(["fecha_exportacion", datetime.now().isoformat(timespec="seconds")])
    meta.append(["nota", "Estimación preliminar. Requiere validación ecológica, económica y jurídica."])
    _ajustar_hoja(meta)

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
